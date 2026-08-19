import os
import json
from datetime import datetime
from openpyxl import Workbook

BASE_DIR = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
REPORTS_DIR = os.path.join(BASE_DIR, "data", "reports")

_applied_jobs: list[dict] = []
_failed_jobs: list[dict] = []
_qa_log: list[dict] = []
_seen_questions: set[str] = set()


def record_applied_jobs(jobs: list[dict]) -> None:
    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    for job in jobs:
        _applied_jobs.append({**job, "applied_at": timestamp})


def record_failed_jobs(jobs: list[dict], reason: str) -> None:
    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    for job in jobs:
        _failed_jobs.append({**job, "attempted_at": timestamp, "reason": reason})


def record_qa(question: str, options: list[str] | None, answer: str) -> None:
    key = question.strip().lower()
    if key in _seen_questions:
        return
    _seen_questions.add(key)
    _qa_log.append({
        "question": question,
        "options": ", ".join(options) if options else "",
        "answer": answer,
    })


def save_and_print_report() -> None:
    timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")

    print("\n📋 Naukri AutoAI: Application Report")
    print("=" * 40)
    print(f"Total Jobs Applied: {len(_applied_jobs)}")

    if _applied_jobs:
        for i, job in enumerate(_applied_jobs, start=1):
            print(f"{i}. {job.get('title', 'Unknown')} @ {job.get('company', 'Unknown')}")
    else:
        print("No jobs were applied to in this run.")

    if _qa_log:
        print("\nQuestions Answered On Your Behalf:")
        for qa in _qa_log:
            print(f"- {qa['question']} -> {qa['answer']}")

    if _failed_jobs:
        print(f"\nSkipped / Failed Attempts: {len(_failed_jobs)}")
        for job in _failed_jobs:
            print(f"- {job.get('title', 'Unknown')} @ {job.get('company', 'Unknown')} ({job.get('reason', '')})")

    os.makedirs(REPORTS_DIR, exist_ok=True)

    json_path = os.path.join(REPORTS_DIR, f"report_{timestamp}.json")
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump({
            "generated_at": timestamp,
            "total_applied": len(_applied_jobs),
            "jobs": _applied_jobs,
            "questions_answered": _qa_log,
            "failed_jobs": _failed_jobs,
        }, f, indent=2)

    excel_path = os.path.join(REPORTS_DIR, f"report_{timestamp}.xlsx")
    _write_excel_report(excel_path)

    print(f"\n📋 Naukri AutoAI: Full report saved to:\n  - {json_path}\n  - {excel_path}")
    print("=" * 40)


def _write_excel_report(path: str) -> None:
    wb = Workbook()

    jobs_sheet = wb.active
    jobs_sheet.title = "Applied Jobs"
    jobs_sheet.append(["#", "Title", "Company", "Job URL", "Applied At"])
    for i, job in enumerate(_applied_jobs, start=1):
        jobs_sheet.append([
            i,
            job.get("title", "Unknown"),
            job.get("company", "Unknown"),
            job.get("url", ""),
            job.get("applied_at", ""),
        ])
    _autosize_columns(jobs_sheet)

    qa_sheet = wb.create_sheet("Answers Filled")
    qa_sheet.append(["Question", "Options Shown", "Answer Submitted"])
    for qa in _qa_log:
        qa_sheet.append([qa["question"], qa["options"], qa["answer"]])
    _autosize_columns(qa_sheet)

    failed_sheet = wb.create_sheet("Skipped or Failed")
    failed_sheet.append(["#", "Title", "Company", "Job URL", "Attempted At", "Reason"])
    for i, job in enumerate(_failed_jobs, start=1):
        failed_sheet.append([
            i,
            job.get("title", "Unknown"),
            job.get("company", "Unknown"),
            job.get("url", ""),
            job.get("attempted_at", ""),
            job.get("reason", ""),
        ])
    _autosize_columns(failed_sheet)

    wb.save(path)


def _autosize_columns(sheet) -> None:
    for column_cells in sheet.columns:
        length = max((len(str(cell.value)) for cell in column_cells if cell.value is not None), default=10)
        sheet.column_dimensions[column_cells[0].column_letter].width = min(length + 2, 80)
